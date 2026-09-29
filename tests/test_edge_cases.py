'''Offline edge-case tests: bad uploads, agent guardrails, and the Gemini -> Groq fallback.

Run from the project folder:   python -m unittest discover -s tests -v
No API key and no internet are needed: the LLM is replaced by a scripted fake model.
'''
import io
import logging
import os
import sys
import unittest
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import pypdf
from langchain_core.documents import Document
from langchain_core.language_models.fake_chat_models import GenericFakeChatModel
from langchain_core.messages import AIMessage
from openpyxl import Workbook

import agent as agent_module
from agent import (FALLBACK_RESPONSE, REFUSAL_MESSAGES, AgentError, DocumentAgent, EvidenceBook,
                   GroundedDraft, Claim, Evidence)
from ingestion import (IngestionError, chunk_documents, describe_empty, load_document,
                       safe_filename)
from llm import (Provider, configured_providers, describe_error, is_daily_limit, is_malformed_tool_call,
                 is_quota_error, is_real_key, redact, resolve_gemini_model, resolve_groq_model,
                 retry_delay_seconds)
from retriever import SearchHit

SAMPLES = ROOT / 'sample_documents'
ERROR_FILES = SAMPLES / 'error_test_files'


def setUpModule():
    logging.disable(logging.CRITICAL)  # the code logs expected failures; keep the test output clean


def tearDownModule():
    logging.disable(logging.NOTSET)


def blank_pdf() -> bytes:
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def locked_pdf() -> bytes:
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt('secret')
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def xlsx_bytes(rows: list[list[str]]) -> bytes:
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class UnsupportedFileTests(unittest.TestCase):
    def assertRejected(self, name: str, data: bytes, *fragments: str) -> None:
        with self.assertRaises(IngestionError) as caught:
            load_document(name, data)
        message = str(caught.exception)
        self.assertIn(name.split('/')[-1], message)
        for fragment in fragments:
            self.assertIn(fragment, message)

    def test_word_documents(self):
        self.assertRejected('report.docx', b'PK\x03\x04data', 'Word', 'PDF')
        self.assertRejected('old.doc', b'\xd0\xcf\x11\xe0', 'Word')

    def test_zip_archives(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, 'w') as archive:
            archive.writestr('a.txt', 'hello')
        self.assertRejected('bundle.zip', buffer.getvalue(), 'Extract')

    def test_other_types(self):
        self.assertRejected('legacy.xls', b'\xd0\xcf\x11\xe0', '.xlsx')
        self.assertRejected('slides.pptx', b'PK\x03\x04', 'PDF')
        self.assertRejected('photo.png', b'\x89PNG\r\n', 'OCR')
        self.assertRejected('notes.md', b'# hi', 'not supported')
        self.assertRejected('noextension', b'hello', 'no extension')

    def test_type_is_checked_before_emptiness(self):
        self.assertRejected('empty.docx', b'', 'Word')

    def test_every_message_lists_supported_types(self):
        with self.assertRaises(IngestionError) as caught:
            load_document('x.docx', b'data')
        self.assertIn('PDF, TXT, CSV and XLSX', str(caught.exception))


class EmptyAndCorruptFileTests(unittest.TestCase):
    def test_empty_and_blank_text(self):
        self.assertEqual(load_document('empty.txt', b''), [])
        self.assertEqual(load_document('blank.txt', b'  \n\t '), [])
        self.assertIn('empty', describe_empty('empty.txt', 0))
        self.assertIn('blank', describe_empty('blank.txt', 5))

    def test_scanned_pdf_without_text(self):
        self.assertEqual(load_document('scan.pdf', blank_pdf()), [])
        self.assertIn('scanned', describe_empty('scan.pdf', 900))

    def test_header_only_csv(self):
        self.assertEqual(load_document('h.csv', b'name,team\n'), [])
        self.assertIn('no data rows', describe_empty('h.csv', 10))

    def test_corrupt_pdf(self):
        with self.assertRaises(IngestionError) as caught:
            load_document('bad.pdf', b'%PDF-1.4 not really a pdf')
        self.assertIn('PDF', str(caught.exception))

    def test_password_protected_pdf(self):
        with self.assertRaises(IngestionError) as caught:
            load_document('locked.pdf', locked_pdf())
        self.assertIn('password', str(caught.exception))

    def test_disguised_files(self):
        for name, data in (('docx_as.pdf', b'PK\x03\x04junk'), ('text_as.xlsx', b'plain text'),
                           ('zip_as.xlsx', b'PK\x03\x04junk')):
            with self.assertRaises(IngestionError, msg=name):
                load_document(name, data)

    def test_binary_text_file(self):
        with self.assertRaises(IngestionError):
            load_document('binary.txt', bytes(range(128, 256)) * 30 + b'\x00\x01')
        with self.assertRaises(IngestionError):
            load_document('binary.csv', b'\x00\x01\x02' * 100)

    def test_oversize_file(self):
        with self.assertRaises(IngestionError) as caught:
            load_document('big.txt', b'x' * (20 * 1024 * 1024 + 1))
        self.assertIn('20 MiB', str(caught.exception))

    def test_unsafe_file_names(self):
        docs = load_document('../../evil/path.txt', b'hello world')
        self.assertEqual(docs[0].metadata['source'], 'path.txt')
        self.assertEqual(safe_filename('..\\..\\a<b>.txt'), 'a_b_.txt')
        self.assertEqual(safe_filename('...'), 'document')


class SupportedFileTests(unittest.TestCase):
    def test_text_encodings(self):
        self.assertEqual(load_document('a.txt', 'Café'.encode('utf-8'))[0].page_content, 'Café')
        self.assertEqual(load_document('a.txt', b'\xef\xbb\xbfBOM ok')[0].page_content, 'BOM ok')
        self.assertEqual(load_document('a.txt', 'Zoë'.encode('utf-16'))[0].page_content, 'Zoë')
        rows = load_document('excel.csv', 'name,city\nJosé,Zürich\n'.encode('cp1252'))
        self.assertIn('José', rows[0].page_content)  # Excel's default CSV encoding

    def test_csv_rows_and_metadata(self):
        docs = load_document('staff.csv', b'name,team\nAva,Legal\nNoah,Finance\n')
        self.assertEqual([d.metadata['row'] for d in docs], [2, 3])
        self.assertEqual(docs[0].metadata['format'], 'csv')

    def test_ragged_csv_does_not_crash(self):
        self.assertTrue(load_document('ragged.csv', b'a,b,c\n1,2\n3,4,5,6\n'))

    def test_xlsx_rows(self):
        docs = load_document('t.xlsx', xlsx_bytes([['Name', 'Dept'], ['Ava', 'Legal'], ['', ''], ['Noah', 'IT']]))
        self.assertEqual(len(docs), 2)
        self.assertEqual(docs[0].metadata['sheet'], 'Sheet')

    def test_upper_case_extension(self):
        self.assertEqual(load_document('NOTES.TXT', b'hello')[0].metadata['format'], 'txt')

    def test_chunk_ids_are_stable(self):
        docs = load_document('p.txt', b'Annual leave is 20 days. ' * 80)
        first, second = chunk_documents(docs), chunk_documents(docs)
        self.assertGreater(len(first), 1)
        self.assertEqual([c.metadata['chunk_id'] for c in first], [c.metadata['chunk_id'] for c in second])
        with self.assertRaises(ValueError):
            chunk_documents(docs, chunk_size=100, chunk_overlap=100)


class ShippedFilesTests(unittest.TestCase):
    '''The exact files an evaluator will find in sample_documents/.'''

    def test_hr_samples_load(self):
        expected = {'Employee_Directory.csv': 1000, 'HR_Analytics_Data.xlsx': 1000,
                    'IT_Security_Memo_Sample.txt': 1}
        for name, count in expected.items():
            docs = load_document(name, (SAMPLES / name).read_bytes())
            self.assertEqual(len(docs), count, name)
        pdf_docs = load_document('HR_Policy_Guide.pdf', (SAMPLES / 'HR_Policy_Guide.pdf').read_bytes())
        self.assertGreater(len(pdf_docs), 50)

    def test_error_test_files_behave_as_documented(self):
        rejected = ['fake_report.docx', 'sample_archive.zip', 'old_excel.xls']
        for name in rejected:
            with self.assertRaises(IngestionError, msg=name):
                load_document(name, (ERROR_FILES / name).read_bytes())
        self.assertEqual(load_document('empty.txt', (ERROR_FILES / 'empty.txt').read_bytes()), [])
        for name in ('corrupted.pdf', 'not_really_excel.xlsx'):
            with self.assertRaises(IngestionError, msg=name):
                load_document(name, (ERROR_FILES / name).read_bytes())
        accented = load_document('a.csv', (ERROR_FILES / 'accented_names_windows1252.csv').read_bytes())
        self.assertEqual(len(accented), 3)
        self.assertIn('José García', accented[0].page_content)


# ---------------------------------------------------------------- agent guardrails
class FakeStore:
    '''Stands in for Chroma: only similarity_search_with_score is used by semantic_search.'''

    def __init__(self, pairs):
        self.pairs = pairs
        self.queries = []

    def similarity_search_with_score(self, query, k=5):
        self.queries.append(query)
        return self.pairs[:k]


LEAVE = Document(page_content='Annual leave is 20 days per year.\nManagers approve requests.',
                 metadata={'source': 'policy.txt', 'chunk_id': 'c1', 'format': 'txt'})


def leave_store() -> FakeStore:
    return FakeStore([(LEAVE, 0.2)])  # cosine distance 0.2 -> similarity 0.8


class ScriptedModel(GenericFakeChatModel):
    '''A chat model that replays fixed messages and accepts tool binding like a real one.'''

    def bind_tools(self, tools, **kwargs):
        return self


def scripted_factory(final_args: dict, query: str = 'annual leave', calls: list | None = None):
    def factory(provider):
        if calls is not None:
            calls.append(provider.name)
        messages = [
            AIMessage(content='', tool_calls=[{'name': 'search_documents', 'args': {'query': query}, 'id': 'call_1'}]),
            AIMessage(content='', tool_calls=[{'name': 'GroundedDraft', 'args': final_args, 'id': 'call_2'}]),
        ]
        return ScriptedModel(messages=iter(messages))
    return factory


def failing_factory(error: Exception):
    def factory(provider):
        raise error
    return factory


GEMINI = Provider('gemini', 'Gemini', 'gemini-test', 'g-secret-key-123456')
GROQ = Provider('groq', 'Groq', 'groq-test', 'q-secret-key-654321')
GOOD_ANSWER = {'outcome': 'answered', 'claims': [{
    'text': 'Employees get 20 days of annual leave.',
    'evidence': [{'source_id': 1, 'quote': 'Annual leave is 20 days per year.'}]}]}


class ValidationTests(unittest.TestCase):
    def book(self) -> EvidenceBook:
        book = EvidenceBook()
        book.add('leave', [SearchHit(LEAVE, 0.8)])
        return book

    def draft(self, quote='Annual leave is 20 days per year.', source_id=1, outcome='answered'):
        return GroundedDraft(outcome=outcome, claims=[Claim(text='20 days.', evidence=[
            Evidence(source_id=source_id, quote=quote)])])

    def test_valid_quote_gives_cited_answer(self):
        answer = DocumentAgent._validate(self.draft(), self.book())
        self.assertIn('[1]', answer.text)
        self.assertEqual(answer.sources[0]['metadata']['source'], 'policy.txt')

    def test_quote_may_differ_only_in_whitespace(self):
        answer = DocumentAgent._validate(self.draft('Annual leave is 20 days per year. Managers approve requests.'), self.book())
        self.assertIn('[1]', answer.text)

    def test_invented_quote_is_refused(self):
        answer = DocumentAgent._validate(self.draft('Annual leave is 45 days per year.'), self.book())
        self.assertEqual(answer.text, FALLBACK_RESPONSE)
        self.assertEqual(answer.sources, [])

    def test_unknown_source_is_refused(self):
        self.assertEqual(DocumentAgent._validate(self.draft(source_id=7), self.book()).text, FALLBACK_RESPONSE)

    def test_no_draft_or_no_claims_is_refused(self):
        self.assertEqual(DocumentAgent._validate(None, self.book()).text, FALLBACK_RESPONSE)
        empty = GroundedDraft(outcome='answered', claims=[])
        self.assertEqual(DocumentAgent._validate(empty, self.book()).text, FALLBACK_RESPONSE)

    def test_each_refusal_outcome_has_a_fixed_message(self):
        for outcome, message in REFUSAL_MESSAGES.items():
            draft = GroundedDraft(outcome=outcome, claims=[])
            self.assertEqual(DocumentAgent._validate(draft, self.book()).text, message)

    def test_evidence_book_deduplicates_and_limits(self):
        book = EvidenceBook()
        first = book.add('q1', [SearchHit(LEAVE, 0.8)])
        again = book.add('q2', [SearchHit(LEAVE, 0.7)])
        self.assertEqual(first[0]['source_id'], again[0]['source_id'])
        self.assertEqual(len(book.hits), 1)
        self.assertEqual([s['query'] for s in book.steps], ['q1', 'q2'])


class AgentFlowTests(unittest.TestCase):
    def test_happy_path_uses_the_search_tool(self):
        store = leave_store()
        agent = DocumentAgent([GEMINI], model_factory=scripted_factory(GOOD_ANSWER))
        answer = agent.answer('How many annual leave days do I get?', store)
        self.assertIn('20 days', answer.text)
        self.assertEqual(answer.provider, 'Gemini')
        self.assertFalse(answer.fell_back)
        self.assertEqual(answer.steps, [{'query': 'annual leave', 'passages': 1}])
        self.assertGreaterEqual(len(store.queries), 2)  # local pre-check + the agent's own search

    def test_model_inventing_a_quote_is_caught(self):
        bad = {'outcome': 'answered', 'claims': [{'text': 'Leave is 45 days.',
               'evidence': [{'source_id': 1, 'quote': 'Annual leave is 45 days per year.'}]}]}
        answer = DocumentAgent([GEMINI], model_factory=scripted_factory(bad)).answer('leave?', leave_store())
        self.assertEqual(answer.text, FALLBACK_RESPONSE)

    def test_model_can_decline_with_fixed_message(self):
        final = {'outcome': 'needs_calculation', 'claims': []}
        answer = DocumentAgent([GEMINI], model_factory=scripted_factory(final)).answer('average leave?', leave_store())
        self.assertEqual(answer.text, REFUSAL_MESSAGES['needs_calculation'])

    def test_no_relevant_passage_never_calls_the_llm(self):
        calls = []
        agent = DocumentAgent([GEMINI], model_factory=scripted_factory(GOOD_ANSWER, calls=calls))
        answer = agent.answer('Who won the 1998 world cup?', FakeStore([(LEAVE, 0.95)]))  # similarity 0.05
        self.assertEqual(answer.text, FALLBACK_RESPONSE)
        self.assertEqual(calls, [])

    def test_empty_and_overlong_questions(self):
        calls = []
        agent = DocumentAgent([GEMINI], model_factory=scripted_factory(GOOD_ANSWER, calls=calls))
        self.assertIn('enter a question', agent.answer('   ', leave_store()).text)
        self.assertIn('shorten', agent.answer('x' * 2001, leave_store()).text)
        self.assertEqual(calls, [])

    def test_needs_at_least_one_provider(self):
        with self.assertRaises(ValueError):
            DocumentAgent([])

    def test_falls_back_to_groq_when_gemini_fails(self):
        calls = []
        primary_down = failing_factory(RuntimeError('429 RESOURCE_EXHAUSTED quota g-secret-key-123456'))
        good = scripted_factory(GOOD_ANSWER, calls=calls)
        agent = DocumentAgent([GEMINI, GROQ], model_factory=lambda p: primary_down(p) if p.name == 'gemini' else good(p))
        answer = agent.answer('How many annual leave days?', leave_store())
        self.assertEqual(answer.provider, 'Groq')
        self.assertTrue(answer.fell_back)
        self.assertEqual(calls, ['groq'])
        self.assertIn('20 days', answer.text)

    def test_failed_provider_is_tried_last_for_a_while(self):
        order = []

        def factory(provider):
            order.append(provider.name)
            if provider.name == 'gemini':
                raise RuntimeError('503 unavailable')
            return scripted_factory(GOOD_ANSWER)(provider)

        agent = DocumentAgent([GEMINI, GROQ], model_factory=factory)
        agent.answer('leave?', leave_store())
        agent.answer('leave again?', leave_store())
        self.assertEqual(order, ['gemini', 'groq', 'groq'])  # second question skips the failed primary

    def test_quota_failure_blocks_the_provider_longer_than_an_outage(self):
        def factory(provider):
            raise RuntimeError('429 RESOURCE_EXHAUSTED quota' if provider.name == 'gemini' else '503 unavailable')

        agent = DocumentAgent([GEMINI, GROQ], model_factory=factory)
        with self.assertRaises(AgentError):
            agent.answer('leave?', leave_store())
        gap = {name: until - __import__('time').monotonic() for name, until in agent._blocked_until.items()}
        self.assertGreater(gap['gemini'], agent_module.PROVIDER_COOLDOWN_SECONDS)   # quota: long pause
        self.assertLessEqual(gap['groq'], agent_module.PROVIDER_COOLDOWN_SECONDS)   # outage: short pause

    def test_all_providers_failing_raises_a_safe_error(self):
        boom = failing_factory(RuntimeError('401 API key not valid g-secret-key-123456'))
        agent = DocumentAgent([GEMINI, GROQ], model_factory=boom)
        with self.assertRaises(AgentError) as caught:
            agent.answer('leave?', leave_store())
        message = str(caught.exception)
        self.assertIn('Gemini', message)
        self.assertIn('Groq', message)
        self.assertNotIn('g-secret-key', message)
        self.assertIn('rejected', message)

    def test_search_failure_becomes_agent_error(self):
        class BrokenStore:
            def similarity_search_with_score(self, query, k=5):
                raise RuntimeError('chroma exploded')
        with self.assertRaises(AgentError):
            DocumentAgent([GEMINI], model_factory=scripted_factory(GOOD_ANSWER)).answer('leave?', BrokenStore())


class LlmHelperTests(unittest.TestCase):
    def test_key_detection(self):
        for value in (None, '', '   ', 'your_gemini_api_key_here', 'your_groq_api_key_here'):
            self.assertFalse(is_real_key(value), value)
        self.assertTrue(is_real_key('AIzaSyExampleKey123'))

    def setUp(self):
        self._saved = {name: os.environ.pop(name, None) for name in (
            'GEMINI_MODEL', 'GEMINI_BACKUP_MODEL', 'GROQ_MODEL', 'GROQ_BACKUP_MODEL', 'PRIMARY_LLM',
            'GOOGLE_API_KEY', 'GROQ_API_KEY')}

    def tearDown(self):
        for name, value in self._saved.items():
            os.environ.pop(name, None)
            if value is not None:
                os.environ[name] = value

    def test_model_resolution(self):
        self.assertEqual(resolve_gemini_model(), 'gemini-3.6-flash')
        self.assertEqual(resolve_gemini_model('gemini-1.5-pro'), 'gemini-3.6-flash')  # retired -> default
        self.assertEqual(resolve_gemini_model('gemini-custom'), 'gemini-custom')
        self.assertEqual(resolve_groq_model(), 'qwen/qwen3.8-27b')
        self.assertEqual(resolve_groq_model('some/custom-model'), 'some/custom-model')

    def test_retired_groq_models_are_replaced_by_the_default(self):
        for retired in ('llama-3.3-70b-versatile', 'llama-3.1-8b-instant', 'gemma2-9b-it', 'mixtral-8x7b-32768',
                        'openai/gpt-oss-120b'):
            self.assertEqual(resolve_groq_model(retired), 'qwen/qwen3.8-27b', retired)
            os.environ['GROQ_MODEL'] = retired            # a stale .env entry must not cause "model not found"
            os.environ['GROQ_BACKUP_MODEL'] = retired
            self.assertEqual(resolve_groq_model(), 'qwen/qwen3.8-27b')
            self.assertEqual([p.model for p in configured_providers(None, 'groq-key-value')],
                             ['qwen/qwen3.8-27b', 'openai/gpt-oss-20b'])

    def test_provider_order_and_placeholders(self):
        both = [p.name for p in configured_providers('gemini-key-value', 'groq-key-value')]
        # the two main models first, then each company's backup model (Groq's first)
        self.assertEqual(both, ['gemini', 'groq', 'groq-backup', 'gemini-backup'])
        self.assertEqual([p.name for p in configured_providers('gemini-key-value', None)], ['gemini', 'gemini-backup'])
        os.environ['GROQ_API_KEY'] = 'your_groq_api_key_here'
        self.assertEqual(configured_providers(None, None), [])
        os.environ['GROQ_API_KEY'] = 'real-groq-key-abc'
        self.assertEqual([p.name for p in configured_providers()], ['groq', 'groq-backup'])  # no Gemini key

    def test_backup_models_share_the_key_but_differ_in_model(self):
        providers = {p.name: p for p in configured_providers('gemini-key-value', 'groq-key-value')}
        self.assertEqual(providers['gemini-backup'].api_key, providers['gemini'].api_key)
        self.assertEqual(providers['groq-backup'].api_key, providers['groq'].api_key)
        self.assertNotEqual(providers['groq-backup'].model, providers['groq'].model)
        os.environ['GROQ_BACKUP_MODEL'] = providers['groq'].model     # same model twice is pointless
        self.assertNotIn('groq-backup', [p.name for p in configured_providers('gemini-key-value', 'groq-key-value')])

    def test_primary_llm_setting_puts_groq_first(self):
        os.environ['PRIMARY_LLM'] = 'groq'
        names = [p.name for p in configured_providers('gemini-key-value', 'groq-key-value')]
        self.assertEqual(names, ['groq', 'gemini', 'groq-backup', 'gemini-backup'])

    def test_quota_errors_are_recognised(self):
        self.assertTrue(is_quota_error(RuntimeError('429 RESOURCE_EXHAUSTED: quota exceeded')))
        self.assertFalse(is_quota_error(RuntimeError('503 UNAVAILABLE high demand')))

    def test_provider_repr_hides_key(self):
        self.assertNotIn('secret', repr(GEMINI))

    def test_error_descriptions(self):
        cases = {
            'API key not valid. Please pass a valid API key.': 'rejected',
            '429 RESOURCE_EXHAUSTED': 'quota',
            '404 NOT_FOUND model': 'model name',
            "The model `llama-3.3-70b-versatile` has been decommissioned": 'model name',
            "The model `x` does not exist or you do not have access to it": 'model name',
            'ConnectError: connection refused': 'busy or unreachable',
            '503 UNAVAILABLE. This model is currently experiencing high demand': 'busy or unreachable',
            '504 DEADLINE_EXCEEDED': 'busy or unreachable',
            'something odd': 'unexpected',
        }
        for text, fragment in cases.items():
            self.assertIn(fragment, describe_error(RuntimeError(text)), text)

    def test_redact(self):
        self.assertEqual(redact('key=abcdefgh1234 failed', ['abcdefgh1234']), 'key=*** failed')


def row_doc(number: int, department: str = 'Sales') -> Document:
    return Document(page_content=f'EmployeeID: {number}\nDepartment: {department}',
                    metadata={'source': 'staff.csv', 'format': 'csv', 'row': number + 1, 'chunk_id': f'r{number}'})


def table_store(rows: int = 5) -> FakeStore:
    return FakeStore([(row_doc(n), 0.3) for n in range(1, rows + 1)])  # similarity 0.7 each


class LocalGuardTests(unittest.TestCase):
    '''Deterministic checks that hold even when the LLM is weak or badly behaved.'''

    def agent(self, calls):
        return DocumentAgent([GEMINI], model_factory=scripted_factory(GOOD_ANSWER, calls=calls))

    def test_calculation_question_over_spreadsheet_rows_is_refused_without_an_llm(self):
        for question in ('How many employees work in the Sales department?',
                         'What is the average salary in Engineering?',
                         'Who is the highest paid employee?',
                         'What is the total salary per department?',
                         'List all employees in Finance',
                         'What is the number of employees in HR?'):
            calls = []
            answer = self.agent(calls).answer(question, table_store())
            self.assertEqual(answer.text, REFUSAL_MESSAGES['needs_calculation'], question)
            self.assertEqual(calls, [], question)
            self.assertIn('no LLM was called', answer.note)

    def test_ordinary_lookups_with_similar_words_are_not_refused(self):
        for question in ('What is the phone number of Luna Sanders?',      # "number of" is not a count
                         "What is Max Miller's job title?",               # Max is a first name
                         'What does the Good performance score mean?',
                         'What is the employee number of Mitchell Serrano?'):
            calls = []
            self.agent(calls).answer(question, table_store())
            self.assertEqual(calls, ['gemini'], question)

    def test_the_same_words_are_fine_when_the_evidence_is_not_a_table(self):
        calls = []
        store = FakeStore([(LEAVE, 0.2)])  # a text passage, not a spreadsheet row
        answer = self.agent(calls).answer('How many annual leave days do I get?', store)
        self.assertEqual(calls, ['gemini'])
        self.assertIn('20 days', answer.text)

    def test_lookup_questions_about_a_row_still_reach_the_llm(self):
        calls = []
        self.agent(calls).answer('What is the department of employee 3?', table_store())
        self.assertEqual(calls, ['gemini'])

    def test_no_evidence_refusal_explains_itself(self):
        answer = self.agent([]).answer('Who won the world cup?', FakeStore([(LEAVE, 0.95)]))
        self.assertEqual(answer.text, FALLBACK_RESPONSE)
        self.assertIn('nothing relevant', answer.note)

    def test_answer_built_from_many_rows_is_refused_as_a_hidden_calculation(self):
        book = EvidenceBook()
        for number in range(1, 6):
            book.add('q', [SearchHit(row_doc(number), 0.7)])

        def draft(ids):
            return GroundedDraft(outcome='answered', claims=[Claim(text='The average is 118,535.', evidence=[
                Evidence(source_id=i, quote=f'EmployeeID: {i}') for i in ids])])

        self.assertEqual(DocumentAgent._validate(draft([1, 2, 3, 4, 5]), book).text, REFUSAL_MESSAGES['needs_calculation'])
        self.assertEqual(DocumentAgent._validate(draft([1, 2, 3]), book).text, REFUSAL_MESSAGES['needs_calculation'])
        self.assertIn('[1]', DocumentAgent._validate(draft([1, 2]), book).text)   # up to two rows is fine
        self.assertIn('[1]', DocumentAgent._validate(draft([1]), book).text)

    def test_malformed_tool_call_is_retried_once_before_falling_back(self):
        calls = []
        state = {'failed': False}

        def flaky(provider):
            calls.append(provider.name)
            if not state['failed']:
                state['failed'] = True
                raise RuntimeError("Error code: 400 - tool_use_failed: attempted to call tool 'commentary'")
            return scripted_factory(GOOD_ANSWER)(provider)

        answer = DocumentAgent([GEMINI, GROQ], model_factory=flaky).answer('How many annual leave days?', leave_store())
        self.assertEqual(calls, ['gemini', 'gemini'])   # same provider twice, no jump to Groq
        self.assertEqual(answer.provider, 'Gemini')
        self.assertFalse(answer.fell_back)


class RateLimitTests(unittest.TestCase):
    def test_retry_hints_are_parsed(self):
        cases = {
            'Please try again in 420ms. Need more tokens?': 0.42,
            'Please try again in 1m3.5s.': 63.5,
            'Please retry in 51.020329022s.': 51.02,
            "'retryDelay': '51s'": 51.0,
            'try again in 2h1m': 7260.0,
        }
        for text, seconds in cases.items():
            self.assertAlmostEqual(retry_delay_seconds(RuntimeError(text)), seconds, places=2, msg=text)
        self.assertIsNone(retry_delay_seconds(RuntimeError('503 unavailable')))

    def test_daily_and_malformed_detection(self):
        self.assertTrue(is_daily_limit(RuntimeError("quotaId: 'GenerateRequestsPerDayPerProjectPerModel-FreeTier'")))
        self.assertTrue(is_daily_limit(RuntimeError('Rate limit reached ... on tokens per day (TPD)')))
        self.assertFalse(is_daily_limit(RuntimeError('Rate limit reached ... on tokens per minute (TPM)')))
        self.assertTrue(is_malformed_tool_call(RuntimeError("code: 'tool_use_failed', failed_generation: '{}'")))
        self.assertFalse(is_malformed_tool_call(RuntimeError('503 unavailable')))

    def test_cooldown_policy(self):
        per_minute = RuntimeError('429 Rate limit reached on tokens per minute (TPM). Please try again in 8.4s.')
        daily = RuntimeError('429 RESOURCE_EXHAUSTED GenerateRequestsPerDayPerProjectPerModel Please retry in 51s')
        unknown = RuntimeError('429 too many requests')
        outage = RuntimeError('503 UNAVAILABLE high demand')
        self.assertAlmostEqual(DocumentAgent._cooldown_seconds(per_minute), 10.4, places=1)   # hint + 2 s
        self.assertEqual(DocumentAgent._cooldown_seconds(daily), agent_module.QUOTA_COOLDOWN_SECONDS)
        self.assertEqual(DocumentAgent._cooldown_seconds(unknown), agent_module.UNKNOWN_LIMIT_COOLDOWN_SECONDS)
        self.assertEqual(DocumentAgent._cooldown_seconds(outage), agent_module.PROVIDER_COOLDOWN_SECONDS)

    def test_numbers_inside_other_numbers_are_not_status_codes(self):
        # "Requested 4013" used to be read as HTTP 401 ("API key rejected").
        exc = RuntimeError('Error code: 429 - Rate limit reached ... Limit 8000, Used 6210, Requested 4013.')
        self.assertIn('quota', describe_error(exc))
        self.assertNotIn('rejected', describe_error(RuntimeError('Requested 4013 tokens, limit 0.403')))
        self.assertIn('rejected', describe_error(RuntimeError('Error code: 401 - Invalid API Key')))
        self.assertIn('busy', describe_error(RuntimeError('Error code: 503')))
        self.assertIn('unusable', describe_error(RuntimeError("tool_use_failed: attempted to call tool 'commentary'")))


if __name__ == '__main__':
    unittest.main(verbosity=2)
